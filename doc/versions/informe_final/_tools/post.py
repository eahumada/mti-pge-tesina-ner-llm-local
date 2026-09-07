import re, copy, docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
d=docx.Document('rebuilt_raw.docx')
def split_para(p, parts):
    anchor=p._p
    # el primero se queda en el párrafo original
    for r in list(p.runs): r._r.getparent().remove(r._r)
    def write(par, text):
        for tok in re.split(r'(\*\*[^*]+\*\*)', text):
            if not tok: continue
            if tok.startswith('**') and tok.endswith('**'): par.add_run(tok[2:-2]).bold=True
            else: par.add_run(tok)
    write(p, parts[0])
    for t in parts[1:]:
        new=copy.deepcopy(anchor)
        for r in new.findall(qn('w:r')): new.remove(r)
        anchor.addnext(new); np=Paragraph(new,d); write(np,t); anchor=new
n=0
for p in list(d.paragraphs):
    t=p.text
    if t.count('Hallazgo ')>1 and re.search(r'Hallazgo \d+:', t):
        parts=re.split(r'(?=Hallazgo \d+:)', t)
        parts=[x.strip() for x in parts if x.strip()]
        parts=[re.sub(r'^(Hallazgo \d+:)', r'**\1**', x) for x in parts]
        split_para(p, parts); n+=1
print('párrafos de hallazgos separados:', n)
# referencias cruzadas a tablas: alinear "la Tabla X" con la leyenda que sigue
body=list(d.element.body); fixed=0
for i,el in enumerate(body):
    if el.tag!=qn('w:p'): continue
    p=Paragraph(el,d)
    m=re.search(r'\bTabla (\d+)\b', p.text)
    if not m or p.style.name=='table caption': continue
    # buscar la próxima leyenda
    for el2 in body[i+1:i+4]:
        if el2.tag==qn('w:p'):
            p2=Paragraph(el2,d)
            if p2.style.name=='table caption':
                m2=re.match(r'Tabla (\d+)\.', p2.text.strip())
                if m2 and m2.group(1)!=m.group(1):
                    for r in p.runs:
                        if f'Tabla {m.group(1)}' in r.text:
                            r.text=r.text.replace(f'Tabla {m.group(1)}', f'Tabla {m2.group(1)}'); fixed+=1
                break
print('referencias cruzadas corregidas:', fixed)
d.save('rebuilt.docx')
