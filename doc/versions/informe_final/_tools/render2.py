import re, copy, sys, docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.table import Table
BASE, OUT = sys.argv[1], sys.argv[2]
USABLE = 8838
MAP = {'Título1':['Título1','Title','Heading 1'],'heading1':['heading1','Heading 1'],'heading2':['heading2','Heading 2'],
       'heading3':['heading3','Heading 3'],'p1a':['p1a','Body Text','Normal'],'Normal':['Normal','Body Text'],'table caption':['table caption','Table Caption','Caption'],
       'programcode':['programcode','Source Code','Normal'],'referenceitem':['referenceitem','Body Text','Normal'],
       'author':['author','Author','Normal'],'address':['address','Author','Normal'],'e-mail':['e-mail','Author','Normal'],
       'abstract':['abstract','Abstract','Normal'],'keywords':['keywords','Abstract','Normal'],
       'bullet item':['bullet item','List Paragraph','Body Text','Normal'],'numbered item':['numbered item','List Paragraph','Body Text','Normal']}
d = docx.Document(BASE)
avail = {s.name: s.style_id for s in d.styles}
def sid(name):
    for cand in MAP.get(name,[name]):
        if cand in avail: return avail[cand]
    return avail.get('Normal')
body = d.element.body
sect = body.find(qn('w:sectPr'))
MODEL_TBL = None
for t in d.tables:
    if len(t.columns)>=5 and len(t.rows)>=3: MODEL_TBL=copy.deepcopy(t._tbl); break
if MODEL_TBL is None and d.tables: MODEL_TBL=copy.deepcopy(d.tables[0]._tbl)
for el in list(body):
    if el is not sect: body.remove(el)
def new_p(style_name):
    p=OxmlElement('w:p'); pPr=OxmlElement('w:pPr')
    st=OxmlElement('w:pStyle'); st.set(qn('w:val'), sid(style_name)); pPr.append(st)
    numPr=OxmlElement('w:numPr'); il=OxmlElement('w:ilvl'); il.set(qn('w:val'),'0')
    ni=OxmlElement('w:numId'); ni.set(qn('w:val'),'0'); numPr.append(il); numPr.append(ni); pPr.insert(0,numPr)
    p.append(pPr)
    body.insert(list(body).index(sect), p)
    return Paragraph(p,d)
INLINE = re.compile(r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*[^*\n]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')
def write_runs(p,text):
    text=re.sub(r'<br\s*/?>',' ',text)
    for tok in INLINE.split(text):
        if not tok: continue
        if tok.startswith('***') and tok.endswith('***'): r=p.add_run(tok[3:-3]); r.bold=True; r.italic=True
        elif tok.startswith('**') and tok.endswith('**'): p.add_run(re.sub(r'\*|`','',tok[2:-2])).bold=True
        elif tok.startswith('`') and tok.endswith('`'): p.add_run(tok[1:-1])
        elif tok.startswith('*') and tok.endswith('*') and len(tok)>2: p.add_run(re.sub(r'\*|`','',tok[1:-1])).italic=True
        elif tok.startswith('[') and '](' in tok:
            m=re.match(r'\[([^\]]+)\]\(([^)]+)\)',tok); p.add_run(m.group(1) if m else tok)
        else: p.add_run(tok)
last_heading=''
pend_caption=None
tras_titulo=False
seccion_abs=None
def emit(style,text):
    global last_heading, tras_titulo, seccion_abs
    if style.startswith('heading') or style=='Título1':
        last_heading=re.sub(r'^[\d.]+\s*','',text)
        tras_titulo=True
        seccion_abs=text.strip().upper() if text.strip().upper() in ('RESUMEN','ABSTRACT') else None
    elif style in ('p1a','Normal','abstract'): tras_titulo=False
    p=new_p(style); write_runs(p,text); return p
def emit_table(rows,caption):
    if caption:
        cp=new_p('table caption'); write_runs(cp,caption)
    tbl=copy.deepcopy(MODEL_TBL); body.insert(list(body).index(sect),tbl); t=Table(tbl,d)
    ncols=max(len(r) for r in rows)
    grid=tbl.find(qn('w:tblGrid')); g=grid.findall(qn('w:gridCol'))
    while len(g)>ncols: grid.remove(g.pop())
    while len(g)<ncols: x=copy.deepcopy(g[-1]); grid.append(x); g.append(x)
    trs=tbl.findall(qn('w:tr'))
    for tr in trs[2:]: tbl.remove(tr)
    for tr in tbl.findall(qn('w:tr')):
        tcs=tr.findall(qn('w:tc'))
        while len(tcs)>ncols: tr.remove(tcs.pop())
        while len(tcs)<ncols: x=copy.deepcopy(tcs[-1]); tr.append(x); tcs.append(x)
    rowsx=tbl.findall(qn('w:tr')); hdr=rowsx[0]; tpl=copy.deepcopy(rowsx[1]) if len(rowsx)>1 else copy.deepcopy(rowsx[0])
    for tr in tbl.findall(qn('w:tr'))[1:]: tbl.remove(tr)
    def fill(tr,cells,ball=False):
        for i,tc in enumerate(tr.findall(qn('w:tc'))):
            v=cells[i] if i<len(cells) else ''
            b=ball or (v.startswith('**') and v.endswith('**')); v=re.sub(r'\*\*|`','',v)
            ps=tc.findall(qn('w:p'))
            for e in ps[1:]: tc.remove(e)
            pp=Paragraph(ps[0],d)
            for r in list(pp.runs): r._r.getparent().remove(r._r)
            for tok in re.split(r'(\*[^*]+\*)', v):
                if not tok: continue
                if tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
                    rr = pp.add_run(tok[1:-1]); rr.italic = True; rr.bold = b
                else:
                    pp.add_run(tok).bold = b
    fill(hdr,rows[0],True)
    for c in rows[1:]:
        tr=copy.deepcopy(tpl); tbl.append(tr); fill(tr,c)
    w=[]
    for ci in range(ncols):
        m=1
        for r in rows:
            v=re.sub(r'\*\*|`','',r[ci]) if ci<len(r) else ''
            m=max(m,min(len(v),34),max((len(x) for x in v.split()),default=1))
        w.append(m)
    tot=sum(w); w=[max(700,int(USABLE*x/tot)) for x in w]; sc=USABLE/sum(w); w=[int(x*sc) for x in w]; w[-1]+=USABLE-sum(w)
    for gc,x in zip(grid.findall(qn('w:gridCol')),w): gc.set(qn('w:w'),str(x))
    for tr in tbl.findall(qn('w:tr')):
        for tc,x in zip(tr.findall(qn('w:tc')),w):
            tcPr=tc.find(qn('w:tcPr'))
            if tcPr is None: tcPr=OxmlElement('w:tcPr'); tc.insert(0,tcPr)
            tcW=tcPr.find(qn('w:tcW'))
            if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.insert(0,tcW)
            tcW.set(qn('w:type'),'dxa'); tcW.set(qn('w:w'),str(x))
src=open('fuente.md',encoding='utf-8').read().split('\n')
i=0;n=len(src);c1=0;tabla=0;pend=None
while i<n:
    s=src[i].strip()
    if not s: i+=1; continue
    if s.startswith('# ') and i<5: emit('Título1',s[2:]); i+=1; continue
    m=re.match(r'^(#{2,4})\s+(.*)$',s)
    if m:
        lvl=len(m.group(1)); ti=m.group(2).strip()
        mm=re.match(r'^(\d+)\.\s+(.*)$',ti)
        if lvl==2 and mm: emit('heading1',f'{mm.group(1)} {mm.group(2)}')
        elif lvl==2: emit('heading1',ti)
        else:
            mm2=re.match(r'^(\d+)\.(\d+)(\.(\d+))?\s+(.*)$',ti)
            if lvl==3 and mm2: emit('heading2',f'{mm2.group(1)}.{mm2.group(2)} {mm2.group(5)}')
            elif lvl==3: emit('heading2',ti)
            elif lvl==4 and mm2: emit('heading3',f'{mm2.group(1)}.{mm2.group(2)}.{mm2.group(4) or ""} {mm2.group(5)}'.replace('. ',' ') if not mm2.group(4) else f'{mm2.group(1)}.{mm2.group(2)}.{mm2.group(4)} {mm2.group(5)}')
            else: emit('heading3',ti)
        i+=1; continue
    if s.startswith('```'):
        i+=1; buf=[]
        while i<n and not src[i].strip().startswith('```'): buf.append(src[i]); i+=1
        i+=1; p=new_p('programcode')
        for k,l in enumerate(buf):
            if k: p.add_run().add_break()
            p.add_run(l)
        continue
    mcap = re.match(r'^[_*]{1,3}\s*Tabla\s+\d+\.\s*(.+?)\s*[_*]{1,3}$', s)
    if mcap:
        pend_caption = mcap.group(1).strip(); i += 1; continue
    if s.startswith('|'):
        rows=[]
        while i<n and src[i].strip().startswith('|'):
            r=src[i].strip()
            if not set(r.replace('|','').strip())<=set(':- '): rows.append([c.strip(' \t') for c in r.strip('|').split('|')])
            i+=1
        tabla+=1
        cap = pend_caption if pend_caption else last_heading
        emit_table(rows,f'Tabla {tabla}. {cap}'); pend_caption=None; tras_titulo=False; continue
    if s.startswith('>'):
        buf=[]
        while i<n and src[i].strip().startswith('>'): buf.append(re.sub(r'^>\s?','',src[i].strip())); i+=1
        t=' '.join(x for x in buf if x.strip())
        if t: emit('p1a',t)
        continue
    if re.match(r'^([-*]|\d+\.)\s+',s):
        while i<n and re.match(r'^\s*([-*]|\d+\.)\s+',src[i].strip()):
            it=re.sub(r'^\s*([-*]|\d+\.)\s+','',src[i].strip())
            emit('bullet item' if re.match(r'^[-*]\s',src[i].strip()) else 'numbered item', it); i+=1
        tras_titulo=False
        continue
    buf=[s]; i+=1
    while i<n and src[i].strip() and not re.match(r'^(#{1,4}\s|\||>|```|[-*]\s|\d+\.\s)',src[i].strip()):
        buf.append(src[i].strip()); i+=1
    txt=' '.join(buf); st='p1a' if tras_titulo else 'Normal'
    if txt.startswith('**Eduardo Mauricio'): st='author'
    elif txt.startswith('Austranet — Departamento'): st='address'
    elif txt.startswith('eahumada@'): st='e-mail'
    elif txt.startswith('**Palabras clave') or txt.startswith('**Keywords'): st='keywords'
    elif re.match(r'^\[\d+\]',txt): st='referenceitem'
    elif seccion_abs in ('RESUMEN','ABSTRACT'): st='abstract'
    emit(st,txt)
d.save(OUT); print(OUT,'->',len(d.paragraphs),'párrafos ·',len(d.tables),'tablas')
