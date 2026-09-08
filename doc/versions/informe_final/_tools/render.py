import re, copy, pickle, docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.text.paragraph import Paragraph
from docx.table import Table

USABLE = 8838
src = open('fuente.md', encoding='utf-8').read().split('\n')
anexos = pickle.load(open('anexos.pkl','rb'))
d = docx.Document('empty.docx')
body = d.element.body
sect = body.find(qn('w:sectPr'))
styles = {s.name: s.style_id for s in d.styles}

def new_p(style_name):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    st = OxmlElement('w:pStyle'); st.set(qn('w:val'), styles.get(style_name, styles['p1a'])); pPr.append(st)
    numPr = OxmlElement('w:numPr')
    il = OxmlElement('w:ilvl'); il.set(qn('w:val'),'0'); numId = OxmlElement('w:numId'); numId.set(qn('w:val'),'0')
    numPr.append(il); numPr.append(numId); pPr.insert(0, numPr)
    p.append(pPr)
    body.insert(list(body).index(sect), p)
    return Paragraph(p, d)

INLINE = re.compile(r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*[^*\n]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')
def write_runs(p, text):
    text = re.sub(r'<br\s*/?>', ' ', text)
    for tok in INLINE.split(text):
        if not tok: continue
        if tok.startswith('***') and tok.endswith('***'):
            r=p.add_run(tok[3:-3]); r.bold=True; r.italic=True
        elif tok.startswith('**') and tok.endswith('**'):
            p.add_run(re.sub(r'\*|`','',tok[2:-2])).bold=True
        elif tok.startswith('`') and tok.endswith('`'):
            p.add_run(tok[1:-1])
        elif tok.startswith('*') and tok.endswith('*') and len(tok)>2:
            p.add_run(re.sub(r'\*|`','',tok[1:-1])).italic=True
        elif tok.startswith('[') and '](' in tok:
            m=re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok)
            p.add_run(m.group(1) if m else tok)
        else:
            p.add_run(tok)

def emit(style, text):
    global last_heading, tras_titulo, seccion_abs
    if style in ('heading1','heading2','heading3','Título1'):
        last_heading = re.sub(r'^[\d.]+\s*','',text)
        tras_titulo = True
        seccion_abs = text.strip().upper() if text.strip().upper() in ('RESUMEN','ABSTRACT') else None
    elif style in ('p1a','Normal','abstract'):
        tras_titulo = False
    p = new_p(style); write_runs(p, text); return p

def emit_table(rows, caption):
    if caption:
        cp = new_p('table caption'); write_runs(cp, caption)
    tbl_el = copy.deepcopy(MODEL_TBL)
    body.insert(list(body).index(sect), tbl_el)
    t = Table(tbl_el, d)
    ncols = max(len(r) for r in rows)
    grid = tbl_el.find(qn('w:tblGrid'))
    gcols = grid.findall(qn('w:gridCol'))
    while len(gcols) > ncols:
        grid.remove(gcols.pop())
    while len(gcols) < ncols:
        g = copy.deepcopy(gcols[-1]); grid.append(g); gcols.append(g)
    # ajustar filas de la plantilla a 1 encabezado + 1 modelo
    trs = tbl_el.findall(qn('w:tr'))
    for tr in trs[2:]: tbl_el.remove(tr)
    for tr in tbl_el.findall(qn('w:tr')):
        tcs = tr.findall(qn('w:tc'))
        while len(tcs) > ncols: tr.remove(tcs.pop())
        while len(tcs) < ncols:
            tc = copy.deepcopy(tcs[-1]); tr.append(tc); tcs.append(tc)
    tpl_row = copy.deepcopy(tbl_el.findall(qn('w:tr'))[1])
    hdr_row = tbl_el.findall(qn('w:tr'))[0]
    for tr in tbl_el.findall(qn('w:tr'))[1:]: tbl_el.remove(tr)
    def fill(tr, cells, bold_all=False):
        for i, tc in enumerate(tr.findall(qn('w:tc'))):
            val = cells[i] if i < len(cells) else ''
            b = bold_all or (val.startswith('**') and val.endswith('**'))
            val = re.sub(r'\*\*|`','', val)
            ps = tc.findall(qn('w:p'))
            for extra in ps[1:]: tc.remove(extra)
            pp = Paragraph(ps[0], d)
            for r in list(pp.runs): r._r.getparent().remove(r._r)
            for tok in re.split(r'(\*[^*]+\*)', val):
                if not tok: continue
                if tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
                    rr = pp.add_run(tok[1:-1]); rr.italic = True; rr.bold = b
                else:
                    pp.add_run(tok).bold = b
    fill(hdr_row, rows[0], bold_all=True)
    for cells in rows[1:]:
        tr = copy.deepcopy(tpl_row); tbl_el.append(tr); fill(tr, cells)
    # anchos proporcionales
    widths = []
    for ci in range(ncols):
        m = 1
        for r in rows:
            v = re.sub(r'\*\*|`','', r[ci]) if ci < len(r) else ''
            m = max(m, min(len(v), 34), max((len(w) for w in v.split()), default=1))
        widths.append(m)
    tot = sum(widths); widths = [max(700, int(USABLE*w/tot)) for w in widths]
    sc = USABLE/sum(widths); widths = [int(w*sc) for w in widths]; widths[-1] += USABLE - sum(widths)
    # suelo por columna DESPUÉS de normalizar: una palabra corta en negrita no debe partirse
    MINW = 900
    falta = sum(max(0, MINW-w) for w in widths)
    if falta:
        widths = [max(MINW, w) for w in widths]
        holgura = [(w-MINW, i) for i, w in enumerate(widths)]
        holgura.sort(reverse=True)
        for _, i in holgura:
            if falta <= 0: break
            quita = min(falta, widths[i]-MINW)
            widths[i] -= quita; falta -= quita
        widths[max(range(len(widths)), key=lambda i: widths[i])] += USABLE - sum(widths)
    for g, w in zip(grid.findall(qn('w:gridCol')), widths): g.set(qn('w:w'), str(w))
    for tr in tbl_el.findall(qn('w:tr')):
        for tc, w in zip(tr.findall(qn('w:tc')), widths):
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None:
                tcPr = OxmlElement('w:tcPr'); tc.insert(0, tcPr)
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW'); tcPr.insert(0, tcW)
            tcW.set(qn('w:type'),'dxa'); tcW.set(qn('w:w'), str(w))
    return t

# modelo de tabla
base = docx.Document('base.docx')
MODEL_TBL = None
for t in base.tables:
    if len(t.columns) >= 5 and len(t.rows) >= 3:
        MODEL_TBL = copy.deepcopy(t._tbl); break

# ---------------- parser ----------------
i = 0; n = len(src)
c1 = c2 = c3 = 0
tabla_n = 0
pend_caption = None
last_heading = ''
tras_titulo = False
seccion_abs = None
anexo_letter = None
log = []
while i < n:
    line = src[i].rstrip()
    s = line.strip()
    if not s:
        i += 1; continue
    # título y cabecera
    if s.startswith('# ') and i < 5:
        emit('Título1', s[2:]); i += 1; continue
    # encabezados
    m = re.match(r'^(#{2,4})\s+(.*)$', s)
    if m:
        lvl = len(m.group(1)); title = m.group(2).strip()
        if re.match(r'^(RESUMEN|ABSTRACT|ÍNDICE DE CONTENIDOS)$', title):
            emit('heading1' if title!='ÍNDICE DE CONTENIDOS' else 'heading1', title)
            i += 1; continue
        mm = re.match(r'^(\d+)\.\s+(.*)$', title)
        if lvl == 2 and mm:
            c1 = int(mm.group(1)); c2 = c3 = 0
            emit('heading1', f'{c1} {mm.group(2)}'); i += 1; continue
        if lvl == 2:
            emit('heading1', title); i += 1; continue
        mm = re.match(r'^(\d+)\.(\d+)\s+(.*)$', title)
        if lvl == 3 and mm:
            c2 = int(mm.group(2)); c3 = 0
            emit('heading2', f'{mm.group(1)}.{mm.group(2)} {mm.group(3)}'); i += 1; continue
        if lvl == 3 and title.startswith('Anexo'):
            emit('heading2', title); anexo_letter = title.split()[1]; i += 1; continue
        if lvl == 3:
            emit('heading2', title); i += 1; continue
        mm = re.match(r'^(\d+)\.(\d+)\.(\d+)\s+(.*)$', title)
        if lvl == 4 and mm:
            emit('heading3', f'{mm.group(1)}.{mm.group(2)}.{mm.group(3)} {mm.group(4)}'); i += 1; continue
        if lvl == 4:
            emit('heading3', title); i += 1; continue
    # bloque de código
    if s.startswith('```'):
        i += 1; buf=[]
        while i < n and not src[i].strip().startswith('```'):
            buf.append(src[i]); i += 1
        i += 1
        p = new_p('programcode')
        for k, l in enumerate(buf):
            if k: p.add_run().add_break()
            p.add_run(l)
        tras_titulo = False
        continue
    # leyenda de tabla escrita en el .md (cursiva): '_Tabla N. Titulo_'
    mcap = re.match(r'^[_*]{1,3}\s*Tabla\s+\d+\.\s*(.+?)\s*[_*]{1,3}$', s)
    if mcap:
        pend_caption = mcap.group(1).strip(); i += 1; continue
    # tabla
    if s.startswith('|'):
        rows=[]
        while i < n and src[i].strip().startswith('|'):
            row = src[i].strip()
            if not set(row.replace('|','').strip()) <= set(':- '):
                rows.append([c.strip(' \t') for c in row.strip('|').split('|')])
            i += 1
        tabla_n += 1
        cap = f'Tabla {tabla_n}. {pend_caption}' if pend_caption else f'Tabla {tabla_n}. {last_heading}'
        emit_table(rows, cap)
        pend_caption = None; tras_titulo = False
        continue
    # cita / nota
    if s.startswith('>'):
        buf=[]
        while i < n and src[i].strip().startswith('>'):
            buf.append(re.sub(r'^>\s?','',src[i].strip())); i += 1
        txt=' '.join(x for x in buf if x.strip())
        if txt: emit('p1a', txt)
        continue
    # listas
    if re.match(r'^([-*]|\d+\.)\s+', s):
        while i < n and re.match(r'^\s*([-*]|\d+\.)\s+', src[i].strip()):
            item = re.sub(r'^\s*([-*]|\d+\.)\s+','', src[i].strip())
            emit('bullet item' if re.match(r'^[-*]\s', src[i].strip()) else 'numbered item', item)
            i += 1
        tras_titulo = False
        continue
    # párrafo normal (unir líneas contiguas)
    buf=[s]; i += 1
    while i < n and src[i].strip() and not re.match(r'^(#{1,4}\s|\||>|```|[-*]\s|\d+\.\s)', src[i].strip()):
        buf.append(src[i].strip()); i += 1
    txt=' '.join(buf)
    style = 'p1a' if tras_titulo else 'Normal'
    if txt.startswith('**Eduardo Mauricio'): style='author'
    elif txt.startswith('Austranet — Departamento'): style='address'
    elif txt.startswith('eahumada@'): style='e-mail'
    elif txt.startswith('**Palabras clave') or txt.startswith('**Keywords'): style='keywords'
    elif re.match(r'^\[\d+\]', txt): style='referenceitem'
    elif seccion_abs in ('RESUMEN','ABSTRACT'): style='abstract'
    emit(style, txt)

d.save('rebuilt_raw.docx')
print('cuerpo reconstruido:', len(d.paragraphs), 'párrafos ·', len(d.tables), 'tablas')
